"""Read paragraphs/tables with explicit Word numbering XML and source locators."""
from collections import defaultdict
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

def read_docx_units(path):
    doc=Document(path)
    try:numbering=doc.part.numbering_part.element
    except (KeyError,NotImplementedError):numbering=None
    definitions={int(a.get(qn('w:abstractNumId'))):a for a in numbering.findall(qn('w:abstractNum'))} if numbering is not None else {}
    numbers={int(n.get(qn('w:numId'))):n for n in numbering.findall(qn('w:num'))} if numbering is not None else {}
    counts={}
    def level_definition(num,abstract,level):
        override=next((x for x in num.findall(qn('w:lvlOverride')) if int(x.get(qn('w:ilvl')))==level),None)
        custom=override.find(qn('w:lvl')) if override is not None else None
        lvl=custom if custom is not None else next((x for x in abstract.findall(qn('w:lvl')) if int(x.get(qn('w:ilvl')))==level),None)
        return lvl,override
    def initial_value(lvl,override):
        changed=override.find(qn('w:startOverride')) if override is not None else None
        start=lvl.find(qn('w:start'))
        value=changed if changed is not None else start
        return int(value.get(qn('w:val'))) if value is not None else 1
    def display_value(value,fmt):
        if fmt=='decimal':return str(value)
        if fmt not in ('upperLetter','lowerLetter'):return None
        display='';cursor=value
        while cursor>0:cursor,rem=divmod(cursor-1,26);display=chr((65 if fmt=='upperLetter' else 97)+rem)+display
        return display
    def paragraph(par,locator):
        text=par.text;pr=par._p.pPr
        numpr=pr.numPr if pr is not None else None
        if numpr is None:
            style=par.style
            for _ in range(5):
                if style is None:break
                spr=style.element.pPr
                if spr is not None and spr.numPr is not None:numpr=spr.numPr;break
                style=style.base_style
        if numpr is None or numpr.numId is None:return text,locator
        numid=numpr.numId.val;level=numpr.ilvl.val if numpr.ilvl is not None else 0;num=numbers.get(numid)
        if num is None:return text,locator
        aid=int(num.find(qn('w:abstractNumId')).get(qn('w:val')));abstract=definitions.get(aid)
        if abstract is None:return text,locator
        lvl,override=level_definition(num,abstract,level)
        if lvl is None:return text,locator
        # OOXML lvlRestart is one-based; omitted means restart on the preceding
        # level (or an earlier ancestor), while zero explicitly disables restart.
        # https://learn.microsoft.com/en-us/dotnet/api/documentformat.openxml.wordprocessing.levelrestart
        for deeper in range(level+1,9):
            dl,do=level_definition(num,abstract,deeper)
            if dl is None:continue
            restart=dl.find(qn('w:lvlRestart'));value=int(restart.get(qn('w:val'))) if restart is not None else deeper
            if value>deeper:value=deeper # invalid value is ignored, default applies
            if value and level<=value-1:counts.pop((numid,deeper),None)
        initial=initial_value(lvl,override)
        key=(numid,level);counts[key]=counts.get(key,initial-1)+1;value=counts[key]
        fmt=lvl.find(qn('w:numFmt'));template=lvl.find(qn('w:lvlText'))
        fmt=fmt.get(qn('w:val')) if fmt is not None else None
        if fmt not in ('decimal','upperLetter','lowerLetter') or template is None:return text,locator
        label=template.get(qn('w:val'))
        for parent in range(level+1):
            pl,po=level_definition(num,abstract,parent)
            pf=pl.find(qn('w:numFmt')) if pl is not None else None
            display=display_value(counts.get((numid,parent),initial_value(pl,po)),pf.get(qn('w:val')) if pf is not None else 'decimal') if pl is not None else None
            if display is not None:label=label.replace('%'+str(parent+1),display)
        if '%' in label:return text,locator
        locator=dict(locator,numbering={'num_id':numid,'abstract_num_id':aid,'level':level,'format':fmt,'value':value,'label':label,'method':'original_word_numbering_xml'})
        return label+' '+text,locator
    units=[];seen_boxes=set()
    def emit_paragraph(child,locator):
        units.append(paragraph(Paragraph(child,doc),locator))
        # Floating Word text boxes can contain the entire A-O bank as a table.
        # Paragraph.text and Document.tables omit these nested containers.
        for bi,box in enumerate(child.xpath('.//w:txbxContent'),1):
            if box in seen_boxes:continue
            seen_boxes.add(box);emit_container(box,dict(locator,text_box=bi))
    def emit_container(container,base):
      for i,child in enumerate(container.iterchildren(),1):
        loc=dict(base,body_block=i) if not base else dict(base,text_box_block=i)
        if child.tag==qn('w:p'):emit_paragraph(child,dict(loc,kind='paragraph'))
        elif child.tag==qn('w:tbl'):
            seen=set()
            for j,row in enumerate(Table(child,doc).rows,1):
                for k,cell in enumerate(row.cells,1):
                    # Retain the element itself: lxml wrapper ids can be reused
                    # after a row is released, causing unrelated cells to vanish.
                    ident=cell._tc
                    if ident in seen:continue
                    seen.add(ident)
                    for pi,par in enumerate(cell.paragraphs,1):emit_paragraph(par._p,dict(loc,table_row=j,table_cell=k,cell_paragraph=pi,kind='table'))
    emit_container(doc.element.body,{})
    return units
